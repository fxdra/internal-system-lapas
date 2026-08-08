import cv2
import sys
import os
import re
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
import pytesseract

img_path = sys.argv[1]
excel_path = sys.argv[2]

image = cv2.imread(img_path)
if image is None:
    print("ERROR|Gagal load image")
    sys.exit(1)

h, w = image.shape[:2]

# =========================
# FORCE CREATE EXCEL
# =========================
if not os.path.exists(excel_path) or os.path.getsize(excel_path) == 0:
    wb = Workbook()
    sheet = wb.active
    sheet['A1'] = 'Nama'
    sheet['B1'] = 'No Reg'
    sheet['C1'] = 'Foto'
    wb.save(excel_path)

# =========================
# LOAD EXCEL
# =========================
wb = load_workbook(excel_path)
sheet = wb.active
row_excel = sheet.max_row + 1

# =========================
# OCR GLOBAL
# =========================
text = pytesseract.image_to_string(image).upper()

nama_match = re.search(r'([A-Z ]+ BIN [A-Z ]+)', text)
reg_match  = re.search(r'[A-Z]{1,3}\s*\d{1,4}\/\d{2}', text)

nama = nama_match.group(0) if nama_match else '-'
reg  = reg_match.group(0) if reg_match else '-'

# =========================
# FACE DETECT (HAAR)
# =========================
gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
face_cascade = cv2.CascadeClassifier(
    cv2.data.haarcascades + 'haarcascade_frontalface_default.xml'
)

faces = face_cascade.detectMultiScale(gray, 1.1, 5)

found = False

for i, (x, y, fw, fh) in enumerate(faces):

    found = True

    # 🔥 framing portrait (ambil kepala aja)
    pad_x = int(fw * 0.6)
    pad_y = int(fh * 1.2)

    x1 = max(0, x - pad_x)
    y1 = max(0, y - pad_y)
    x2 = min(w, x + fw + pad_x)
    y2 = min(h, y + fh + int(fh * 0.5))

    crop = image[y1:y2, x1:x2]

    save_path = img_path.replace('.jpg', f'_face_{i}.jpg')
    cv2.imwrite(save_path, crop)

    # EXCEL
    sheet[f'A{row_excel}'] = nama
    sheet[f'B{row_excel}'] = reg

    img_excel = XLImage(save_path)
    img_excel.width = 90
    img_excel.height = 90
    sheet.add_image(img_excel, f'C{row_excel}')

    print(f"{nama}|{reg}|{save_path}")

    row_excel += 1

# =========================
# FALLBACK
# =========================
if not found:

    crop = image[int(h*0.2):int(h*0.8), 0:int(w*0.4)]

    save_path = img_path.replace('.jpg', '_fallback.jpg')
    cv2.imwrite(save_path, crop)

    sheet[f'A{row_excel}'] = nama
    sheet[f'B{row_excel}'] = reg

    img_excel = XLImage(save_path)
    img_excel.width = 90
    img_excel.height = 90
    sheet.add_image(img_excel, f'C{row_excel}')

    print(f"{nama}|{reg}|{save_path}")

    row_excel += 1

# =========================
# SAVE
# =========================
wb.save(excel_path)